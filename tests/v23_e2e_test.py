#!/usr/bin/env python3
"""Run the installed V2.3 Profiles and the public assignment migration chain.

This forward test deliberately stops before human visual approval.  CI may prove
that the automated compile gate passes and that finalization remains blocked; it
must never manufacture a ``visual-review.json`` with a passed status.
"""
from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import sys
from pathlib import Path
from typing import Any

from docx import Document

REPOSITORY = Path(__file__).resolve().parents[1]
SCRIPT_DIR = REPOSITORY / "skills" / "make-bilingual-study-pdf" / "scripts"
sys.path.insert(0, str(SCRIPT_DIR))
sys.path.insert(0, str(REPOSITORY / "tests"))

import docx_style
from common import read_json, read_jsonl, sha256_file, write_json, write_jsonl
from docx_ast import transform
from pipeline import report_status
from v2_assignment_chain_diff_test import run_assignment_document_chain, finish_assignment_document_chain, _run_pipeline_stage, _docx_projection, _render_projection
from v2_migration_contract_test import build_candidate_v2
from profile import load_profile


FIXTURE_ROOT = REPOSITORY / "tests" / "fixtures" / "mineru" / "pipeline-3.4.4"
PROFILES = ("academic-paper-en-zh", "lecture-notes-en-zh")


def run_script(script: str, *arguments: str, check: bool = True) -> subprocess.CompletedProcess[str]:
    command = [sys.executable, str(SCRIPT_DIR / script), *arguments]
    process = subprocess.run(command, check=False, text=True)
    if check and process.returncode:
        raise RuntimeError(f"command failed ({process.returncode}): {command!r}")
    return process


def require_tool(*names: str) -> str:
    for name in names:
        resolved = shutil.which(name)
        if resolved:
            return resolved
    raise RuntimeError(f"required executable is unavailable: {' or '.join(names)}")


def tool_version(command: list[str]) -> str:
    process = subprocess.run(command, check=False, capture_output=True, text=True)
    output = (process.stdout + "\n" + process.stderr).strip()
    return output.splitlines()[0] if output else f"exit={process.returncode}"


def synthesize_responses(work_dir: Path) -> int:
    """Fill every frozen request with deterministic, low-copy synthetic Chinese."""
    plan = read_json(work_dir / "translation" / "plan.json")
    total = 0
    for batch in plan["batches"]:
        requests = read_jsonl(work_dir / "translation" / batch["request_file"])
        responses = []
        for request in requests:
            placeholders = [
                item["placeholder"] for item in request.get("protected_tokens", [])
            ]
            translation = (
                "译文：这是用于验证完整英中双语生成链、角色处置和冻结证据的"
                "原创合成中文内容。"
            )
            if placeholders:
                translation += " 保留项：" + " ".join(placeholders)
            responses.append(
                {
                    "id": request["id"],
                    "source_sha256": request["source_sha256"],
                    "translation": translation,
                }
            )
        response_path = work_dir / "translation" / batch["response_file"]
        write_jsonl(response_path, responses)
        total += len(responses)
    if total != plan["expected_segment_count"]:
        raise RuntimeError("synthesized response count does not match frozen plan")
    return total


def assert_automated_forward(profile: str, output_root: Path) -> dict[str, Any]:
    work_dir = output_root / profile
    basename = f"{profile}-forward"
    source = FIXTURE_ROOT / "native-source.pdf"
    mineru_output = FIXTURE_ROOT / "native"

    run_script(
        "pipeline.py",
        "import-mineru",
        str(source),
        str(mineru_output),
        "--work-dir",
        str(work_dir),
        "--profile",
        profile,
    )
    source_audit = read_json(work_dir / "source-audit.json")
    if source_audit.get("status") != "passed":
        raise RuntimeError(f"{profile} source audit did not pass")
    ir = read_json(work_dir / "document-ir.json")
    counts = ir["inventories"]["semantic_role_counts"]
    required_counts = (
        {
            "title": 1,
            "author-affiliation": 1,
            "abstract": 1,
            "section": 1,
            "paragraph": 1,
            "references": 1,
        }
        if profile == "academic-paper-en-zh"
        else {
            "title": 1,
            "section": 1,
            "paragraph": 1,
            "definition": 1,
            "theorem": 1,
            "proof": 1,
            "example": 1,
        }
    )
    missing_roles = {
        role: {"minimum": minimum, "actual": counts.get(role, 0)}
        for role, minimum in required_counts.items()
        if counts.get(role, 0) < minimum
    }
    if missing_roles:
        raise RuntimeError(f"{profile} fixture role coverage is incomplete: {missing_roles}")

    run_script("pipeline.py", "prepare", str(work_dir))
    translated_segments = synthesize_responses(work_dir)
    run_script("audit_translation.py", str(work_dir))
    run_script("pipeline.py", "build", str(work_dir), "--basename", basename)
    markdown = work_dir / "output" / f"{basename}.md"
    run_script(
        "pipeline.py",
        "docx",
        str(work_dir),
        "--markdown",
        str(markdown),
        "--basename",
        basename,
        "--minimum-images",
        "1",
    )
    run_script(
        "pipeline.py",
        "compile-docx",
        str(work_dir),
        "--basename",
        basename,
        "--dpi",
        "96",
    )

    compile_audit = read_json(work_dir / "output" / "compile-audit.json")
    if compile_audit.get("automated_status") != "passed":
        raise RuntimeError(f"{profile} automated compile audit did not pass")
    contact_sheets = compile_audit.get("contact_sheets")
    if not isinstance(contact_sheets, list) or not contact_sheets:
        raise RuntimeError(f"{profile} compile audit has no contact-sheet evidence")

    finalization = run_script("pipeline.py", "finalize", str(work_dir), check=False)
    if finalization.returncode == 0:
        raise RuntimeError("CI forward test unexpectedly finalized without human review")
    status = report_status(work_dir)
    if status["gate_statuses"].get("visual_review") == "passed":
        raise RuntimeError("CI forward test must not create a passed visual review")
    if status["next_action"] != "inspect every final render and record the visual review":
        raise RuntimeError(f"unexpected post-compile next action: {status['next_action']}")

    build = read_json(work_dir / "output" / "build-manifest.json")
    pdf = work_dir / "output" / f"{basename}.pdf"
    docx = work_dir / "output" / f"{basename}.docx"
    return {
        "profile": profile,
        "work_dir": str(work_dir),
        "translated_segments": translated_segments,
        "role_inventory": ir["inventories"]["role_inventory"],
        "source_audit_sha256": sha256_file(work_dir / "source-audit.json"),
        "build_manifest_sha256": sha256_file(work_dir / "output" / "build-manifest.json"),
        "docx_sha256": sha256_file(docx),
        "pdf_sha256": sha256_file(pdf),
        "page_count": compile_audit["page_count"],
        "render_count": compile_audit["rendered_page_count"],
        "contact_sheets": contact_sheets,
        "automated_status": compile_audit["automated_status"],
        "finalization_without_visual_review": "blocked",
        "build_markdown": build["markdown"],
    }


def assert_assignment_v1_render(
    output_root: Path, office: str, pdftoppm: str, pdftotext: str
) -> dict[str, Any]:
    """Render a real V1 Problem pair and reject visible STYLEREF errors."""
    probe = output_root / "assignment-v1-render"
    probe.mkdir()
    profile = load_profile("assignment-en-zh")
    converted_ast = subprocess.run(
        ["pandoc", "--from", "markdown", "--to", "json"],
        input=(
            "**Problem (p1): Verify the complete assignment render path.**\n\n"
            "> **问题（p1）：验证完整的作业渲染链。**\n"
        ),
        check=True,
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    grouped = transform(json.loads(converted_ast.stdout), profile)
    problem_count = grouped["meta"]["v2-problem-group-count"]["c"]
    if problem_count != "1":
        raise RuntimeError(f"assignment V1 render grouped {problem_count} Problems")
    grouped_path = probe / "grouped.json"
    raw_docx_path = probe / "raw.docx"
    write_json(grouped_path, grouped)
    subprocess.run(
        [
            "pandoc",
            str(grouped_path),
            "--from",
            "json",
            "--to",
            "docx",
            "--output",
            str(raw_docx_path),
        ],
        check=True,
        capture_output=True,
        text=True,
    )
    document = Document(raw_docx_path)
    docx_style.configure_profile(profile)
    docx_style.apply_styles(
        document,
        document_title="Assignment V1 render probe",
        header_label=profile["render"]["docx"]["header_label"],
        footer_label=profile["render"]["docx"]["footer_label"],
    )
    docx_path = probe / "assignment-v1-render.docx"
    pdf_path = probe / "assignment-v1-render.pdf"
    render_prefix = probe / "assignment-v1-render"
    document.save(docx_path)
    converted = subprocess.run(
        [
            office,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(probe),
            str(docx_path),
        ],
        check=False,
        capture_output=True,
        text=True,
    )
    if converted.returncode or not pdf_path.is_file():
        raise RuntimeError(
            "header fallback LibreOffice conversion failed: "
            f"{converted.stdout}{converted.stderr}"
        )
    extracted = subprocess.run(
        [pdftotext, str(pdf_path), "-"],
        check=True,
        capture_output=True,
        text=True,
    ).stdout
    if "Error: Reference source not found" in extracted:
        raise RuntimeError("header fallback rendered an unresolved STYLEREF field")
    if profile["render"]["docx"]["header_label"] not in extracted:
        raise RuntimeError("header fallback label is absent from rendered PDF text")
    for expected in (
        "Problem (p1): Verify the complete assignment render path.",
        "问题（p1）：验证完整的作业渲染链。",
    ):
        if expected not in extracted:
            raise RuntimeError(f"assignment V1 render lost expected text: {expected}")
    subprocess.run(
        [
            pdftoppm,
            "-png",
            "-r",
            "96",
            "-singlefile",
            str(pdf_path),
            str(render_prefix),
        ],
        check=True,
        capture_output=True,
        text=True,
    )
    render_path = probe / "assignment-v1-render.png"
    if not render_path.is_file():
        raise RuntimeError("header fallback render was not created")
    return {
        "status": "passed",
        "docx_sha256": sha256_file(docx_path),
        "pdf_sha256": sha256_file(pdf_path),
        "render_sha256": sha256_file(render_path),
        "header_label": profile["render"]["docx"]["header_label"],
        "problem_group_count": 1,
        "unresolved_reference_error": False,
    }


def assert_assignment_migration_forward(output_root: Path) -> dict[str, Any]:
    """Exercise real V1 and candidate-V2 DOCX/PDF gates on public input.

    This is intentionally an automated-forward proof only.  It produces real
    contact sheets, verifies the production document gates, and proves final QA
    remains blocked until a human has inspected those images.
    """
    v1_profile = load_profile("assignment-en-zh")
    v1 = run_assignment_document_chain(output_root / "assignment-v1", v1_profile)
    v2 = run_assignment_document_chain(
        output_root / "assignment-v2", build_candidate_v2(v1_profile)
    )
    from migrate_profile_contract_test import migrate_profile_reference, tree_bytes
    migrated_work = output_root / "assignment-existing-migrated" / "work"
    shutil.copytree(Path(v1["work_dir"]), migrated_work)
    backup = output_root / "assignment-migration-backup"
    observed = tree_bytes(migrated_work)
    transaction = migrate_profile_reference(migrated_work, backup, dry_run=True)
    if tree_bytes(migrated_work) != observed:
        raise RuntimeError("migration dry-run changed existing WORK")
    if migrate_profile_reference(migrated_work, backup) != transaction:
        raise RuntimeError("migration differs from its dry-run plan")
    _run_pipeline_stage("source-audit", migrated_work)
    _run_pipeline_stage("prepare", migrated_work, "--max-source-chars", "1000")
    migrated = finish_assignment_document_chain(migrated_work, build_candidate_v2(v1_profile))
    for key in ("projection", "docx_projection", "render_projection"):
        if migrated[key] != v2[key]:
            raise RuntimeError(f"existing WORK migration differs from fresh V2: {key}")
    write_json(output_root / "existing-work-migration.json", {
        "transaction": transaction, "dry_run_zero_write": True,
        "source_audit_rebuilt": "passed", "fresh_v2_equal": True,
        "finalize": migrated["finalize"],
    })
    if v1["projection"] != v2["projection"]:
        raise RuntimeError("assignment V1/V2 normalized visible projection drifted")
    if v1["docx_projection"] != v2["docx_projection"]:
        raise RuntimeError("assignment V1/V2 normalized DOCX structure drifted")
    if v1["render_projection"] != v2["render_projection"]:
        raise RuntimeError("assignment V1/V2 per-page render projection drifted")
    for result in (v1, v2):
        if result["docx_audit"].get("status") != "passed":
            raise RuntimeError("assignment DOCX audit did not pass")
        if result["compile_audit"].get("automated_status") != "passed":
            raise RuntimeError("assignment PDF compile did not pass")
        if not result["compile_audit"].get("contact_sheets"):
            raise RuntimeError("assignment compile lacks visual contact sheets")
        if result["gate_statuses"].get("qa_report") == "passed":
            raise RuntimeError("assignment final QA must remain human-review gated")
        if (
            result["finalize"]["returncode"] == 0
            or not result["finalize"]["gate_hashes_unchanged"]
            or result["finalize"]["qa_passed"]
        ):
            raise RuntimeError(f"assignment finalize fail-closed drifted: {result['finalize']}")
    return {
        "v1": {
            "work_dir": v1["work_dir"],
            "profile_schema_version": v1["profile_schema_version"],
            "docx_audit_sha256": sha256_file(
                Path(v1["work_dir"]) / "output" / "docx-audit.json"
            ),
            "compile_audit_sha256": sha256_file(
                Path(v1["work_dir"]) / "output" / "compile-audit.json"
            ),
            "docx_sha256": v1["docx_sha256"],
            "pdf_sha256": v1["pdf_sha256"],
            "finalize": v1["finalize"],
        },
        "v2": {
            "work_dir": v2["work_dir"],
            "profile_schema_version": v2["profile_schema_version"],
            "docx_audit_sha256": sha256_file(
                Path(v2["work_dir"]) / "output" / "docx-audit.json"
            ),
            "compile_audit_sha256": sha256_file(
                Path(v2["work_dir"]) / "output" / "compile-audit.json"
            ),
            "docx_sha256": v2["docx_sha256"],
            "pdf_sha256": v2["pdf_sha256"],
            "finalize": v2["finalize"],
        },
        "normalized_projection": v1["projection"],
        "normalized_docx_projection": v1["docx_projection"],
        "per_page_render_projection": v1["render_projection"],
        "finalization_without_human_visual_review": "blocked",
    }



def assert_production_base_visible(output_root: Path, toolchain: dict) -> dict:
    baseline = read_json(REPOSITORY / "tests/fixtures/profiles/production-v2-base-visible.json")
    if baseline["toolchain"] != toolchain:
        raise RuntimeError("production base comparison requires the frozen exact toolchain; regenerate both sides together")
    evidence = {}
    for profile, expected in baseline["profiles"].items():
        output = output_root / profile / "output"
        compile_report = read_json(output / "compile-audit.json")
        actual = {"docx": _docx_projection(output / compile_report["docx"]),
                  "render": _render_projection(output / "pdf-renders")}
        if actual != expected:
            raise RuntimeError(f"{profile} DOCX paragraphs or page renders drifted from exact base {baseline['base_sha']}")
        evidence[profile] = actual
    return {"base_sha": baseline["base_sha"], "base_run_id": baseline["run_id"],
            "base_artifact_id": baseline["artifact_id"], "status": "passed", "profiles": evidence}


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Run V2.3 academic-paper and lecture-notes automated forward tests."
    )
    parser.add_argument("output_root", type=Path)
    args = parser.parse_args()

    require_tool("pandoc")
    office = require_tool("libreoffice", "soffice")
    pdftoppm = require_tool("pdftoppm")
    pdftotext = require_tool("pdftotext")
    require_tool("pdffonts")
    require_tool("fc-match")

    output_root = args.output_root.expanduser().resolve()
    if output_root.exists() and any(output_root.iterdir()):
        raise SystemExit(f"output root must be absent or empty: {output_root}")
    output_root.mkdir(parents=True, exist_ok=True)

    results = [assert_automated_forward(profile, output_root) for profile in PROFILES]
    assignment_v1_render = assert_assignment_v1_render(
        output_root, office, pdftoppm, pdftotext
    )
    assignment_migration = assert_assignment_migration_forward(output_root)
    from v2_assignment_chain_diff_test import write_repeated_visual_source
    repeated_source = write_repeated_visual_source(output_root / "repeated-visuals.pdf")
    repeated = run_assignment_document_chain(
        output_root / "assignment-repeated-visuals",
        build_candidate_v2(load_profile("assignment-en-zh")), repeated_source,
    )
    if repeated["docx_audit"]["role_occurrence_evidence"]["math-with-text"]["visual"] != 2:
        raise RuntimeError("repeated visual asset occurrences were lost")
    write_json(output_root / "repeated-visuals-report.json", repeated)
    report = {
        "schema_version": 1,
        "status": "passed",
        "scope": "automated-forward-only-human-visual-review-required",
        "fixture": "tests/fixtures/mineru/pipeline-3.4.4",
        "mineru_version": "3.4.4",
        "mineru_backend": "pipeline",
        "toolchain": {
            "python": sys.version.split()[0],
            "pandoc": tool_version([require_tool("pandoc"), "--version"]),
            "libreoffice": tool_version([office, "--version"]),
            "pdftoppm": tool_version([require_tool("pdftoppm"), "-v"]),
            "fc_match": tool_version([require_tool("fc-match"), "--version"]),
        },
        "profiles": results,
        "assignment_v1_render": assignment_v1_render,
        "assignment_migration": assignment_migration,
    }
    report["production_base_visible"] = assert_production_base_visible(output_root, report["toolchain"])
    write_json(output_root / "v23-e2e-report.json", report)
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
